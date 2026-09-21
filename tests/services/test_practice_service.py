"""Tests for the practice session (FR-5.x, FR-6.x).

The invariant most worth protecting: a wrong answer costs score but never erases a
word's accumulated mastery progress.
"""

from __future__ import annotations

import random
from datetime import date, timedelta

import pytest

from vocabulary_trainer.data.history_store import HistoryStore
from vocabulary_trainer.data.master_file_repository import MasterFileRepository
from vocabulary_trainer.data.practice_log_repository import PracticeLogRepository
from vocabulary_trainer.domain.models import Collection, CollectionSortMode
from vocabulary_trainer.services.clock import FixedClock
from vocabulary_trainer.services.errors import EmptyPoolError
from vocabulary_trainer.services.practice_service import PracticeService

from conftest import FakeAudio, FakeTts, make_entry


@pytest.fixture
def service(
    repository: MasterFileRepository,
    history: HistoryStore,
    fake_tts: FakeTts,
    fake_audio: FakeAudio,
    clock: FixedClock,
) -> PracticeService:
    return PracticeService(
        repository,
        history,
        tts=fake_tts,
        audio=fake_audio,
        clock=clock,
        rng=random.Random(1234),
    )


def seed(
    repository: MasterFileRepository, collection: str, words: list[str]
) -> None:
    existing = {c.name for c in repository.all_collections()}
    if collection not in existing:
        repository.insert_collection(Collection(collection, date(2026, 1, 1)))
    for word in words:
        repository.insert_word(make_entry(word, collection=collection))


class TestCollectionListing:
    def test_lists_collections_with_counts(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha", "beta"])
        summaries = {s.name: s.word_count for s in service.list_collections()}
        assert summaries["IELTS"] == 2

    def test_search_filters_case_insensitively(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS Vocabulary", ["a"])
        seed(repository, "Daily Reading", ["b"])

        names = [s.name for s in service.list_collections(search="ielts")]
        assert names == ["IELTS Vocabulary"]

    def test_search_matches_a_substring_anywhere(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "Business English", ["a"])
        names = [s.name for s in service.list_collections(search="ness eng")]
        assert names == ["Business English"]

    def test_no_matches_yields_an_empty_list(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["a"])
        assert service.list_collections(search="zzzz") == []

    def test_blank_search_returns_everything(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["a"])
        assert len(service.list_collections(search="   ")) == 2  # + General

    def test_newest_first_is_the_default(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        repository.insert_collection(Collection("Older", date(2026, 1, 1)))
        repository.insert_collection(Collection("Newer", date(2026, 8, 1)))
        names = [s.name for s in service.list_collections()]
        assert names.index("Newer") < names.index("Older")

    def test_oldest_first(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        repository.insert_collection(Collection("Older", date(2026, 1, 1)))
        repository.insert_collection(Collection("Newer", date(2026, 8, 1)))
        names = [
            s.name
            for s in service.list_collections(sort=CollectionSortMode.OLDEST_FIRST)
        ]
        assert names.index("Older") < names.index("Newer")

    def test_name_ascending(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        repository.insert_collection(Collection("Zebra", date(2026, 8, 1)))
        repository.insert_collection(Collection("apple", date(2026, 1, 1)))
        names = [
            s.name for s in service.list_collections(sort=CollectionSortMode.NAME_ASC)
        ]
        assert names == ["apple", "General", "Zebra"]

    def test_search_and_sort_compose(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        repository.insert_collection(Collection("Test Beta", date(2026, 1, 1)))
        repository.insert_collection(Collection("Test Alpha", date(2026, 8, 1)))
        repository.insert_collection(Collection("Unrelated", date(2026, 5, 1)))

        names = [
            s.name
            for s in service.list_collections(
                search="test", sort=CollectionSortMode.NAME_ASC
            )
        ]
        assert names == ["Test Alpha", "Test Beta"]

    def test_equal_dates_keep_a_stable_order(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """E5-S2 branch 4a: no shuffling between renders."""
        shared = date(2026, 4, 1)
        repository.insert_collection(Collection("Bravo", shared))
        repository.insert_collection(Collection("Alpha", shared))

        first = [s.name for s in service.list_collections()]
        second = [s.name for s in service.list_collections()]
        assert first == second


class TestPoolSizing:
    def test_combined_count_across_collections(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["a", "b", "c"])
        seed(repository, "Daily", ["d", "e"])
        assert service.combined_word_count(["IELTS", "Daily"]) == 5

    def test_count_matches_case_insensitively(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["a", "b"])
        assert service.combined_word_count(["ielts"]) == 2

    def test_no_selection_is_zero(self, service: PracticeService) -> None:
        assert service.combined_word_count([]) == 0

    def test_can_start_requires_at_least_one_word(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """FR-5.8."""
        assert service.can_start([]) is False
        assert service.can_start(["General"]) is False  # seeded but empty
        seed(repository, "IELTS", ["a"])
        assert service.can_start(["IELTS"]) is True

    def test_only_empty_collections_still_blocks_start(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        repository.insert_collection(Collection("Empty", date(2026, 1, 1)))
        assert service.can_start(["Empty", "General"]) is False

    @pytest.mark.parametrize(
        ("requested", "expected"), [(0, 1), (-3, 1), (1, 1), (3, 3), (10, 10), (11, 10), (99, 10)]
    )
    def test_required_writes_is_clamped(
        self, service: PracticeService, requested: int, expected: int
    ) -> None:
        """FR-5.6: the stepper range is 1 to 10."""
        assert service.clamp_required_writes(requested) == expected


class TestSessionStart:
    def test_builds_a_pool_from_checked_collections(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha", "beta", "gamma"])
        session = service.start_session(["IELTS"], 3)
        assert session.initial_pool_size == 3

    def test_excludes_unchecked_collections(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha", "beta"])
        seed(repository, "Other", ["gamma"])
        session = service.start_session(["IELTS"], 1)
        assert session.initial_pool_size == 2

    def test_empty_pool_is_refused(self, service: PracticeService) -> None:
        with pytest.raises(EmptyPoolError):
            service.start_session(["General"], 3)

    def test_progress_denominator_is_pool_times_threshold(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """FR-6.3."""
        seed(repository, "IELTS", [f"w{i}" for i in range(10)])
        session = service.start_session(["IELTS"], 3)
        assert session.total_required_attempts == 30

    def test_single_word_pool_is_valid(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["lonely"])
        assert service.start_session(["IELTS"], 2).initial_pool_size == 1

    def test_threshold_is_clamped_at_start(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["a"])
        assert service.start_session(["IELTS"], 500).required_correct == 10


class TestCorrectAnswers:
    def test_correct_answer_raises_the_score_and_plays_feedback(
        self,
        service: PracticeService,
        repository: MasterFileRepository,
        fake_audio: FakeAudio,
        fake_tts: FakeTts,
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        session = service.start_session(["IELTS"], 3)
        word = session.current.entry.word  # type: ignore[union-attr]

        outcome = service.submit_attempt(word)

        assert outcome.was_correct is True
        assert outcome.score == 1
        assert outcome.correct_count == 1
        assert outcome.wrong_count == 0
        assert outcome.remaining_repetitions == 2
        assert fake_audio.correct_plays == 1
        assert fake_tts.spoken == [word]

    def test_matching_ignores_case_and_padding(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 3)
        assert service.submit_attempt("  ALPHA  ").was_correct is True

    def test_word_is_mastered_at_the_threshold(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """FR-6.6, clean run: no mistakes, so the target never moves."""
        seed(repository, "IELTS", ["alpha"])
        session = service.start_session(["IELTS"], 2)

        first = service.submit_attempt("alpha")
        assert first.became_mastered is False
        assert first.remaining_repetitions == 1

        second = service.submit_attempt("alpha")
        assert second.became_mastered is True
        assert second.remaining_repetitions == 0
        assert session.is_finished is True

    def test_unmastered_word_returns_to_the_pool(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha", "beta"])
        session = service.start_session(["IELTS"], 3)
        first_word = session.current.entry.word  # type: ignore[union-attr]

        service.submit_attempt(first_word)

        assert session.pool_size == 2
        assert session.current.entry.word != first_word  # type: ignore[union-attr]

    def test_mastering_the_last_word_finishes_the_session(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        session = service.start_session(["IELTS"], 1)
        service.submit_attempt("alpha")
        assert session.is_finished is True
        assert session.words_mastered == 1


class TestWrongAnswers:
    def test_penalty_is_applied_and_score_drops(
        self,
        service: PracticeService,
        repository: MasterFileRepository,
        fake_audio: FakeAudio,
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 3)

        outcome = service.submit_attempt("wrong")

        assert outcome.was_correct is False
        assert outcome.total_penalties == 1
        assert outcome.score == -1
        assert fake_audio.incorrect_plays == 1

    def test_correct_spelling_is_revealed(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """FR-6.8."""
        seed(repository, "IELTS", ["curiosity"])
        service.start_session(["IELTS"], 3)
        assert service.submit_attempt("curiousity").correct_spelling == "curiosity"

    def test_a_miss_never_erases_correct_count_but_does_raise_the_target(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """Replaces the old streak invariant. Under the NOR rule a miss no longer
        leaves mastery progress untouched -- it adds one repetition to the target,
        on top of the score penalty. What is still guaranteed is that
        ``correct_count`` itself is never decremented by a miss."""
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 5)

        service.submit_attempt("alpha")
        service.submit_attempt("alpha")
        state = service.current_word()
        assert state.correct_count == 2  # type: ignore[union-attr]

        outcome = service.submit_attempt("wrong")

        assert outcome.correct_count == 2, "a miss must not erase correct_count"
        assert outcome.wrong_count == 1
        # required(5) + wrong(1) - correct(2) = 4 remaining, one more than before
        # the miss (it was 3 remaining after the second correct answer).
        assert outcome.remaining_repetitions == 4
        assert state.correct_count == 2  # type: ignore[union-attr]

    def test_requested_example_one_wrong_answer_adds_one_repetition(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """NODR=3, one wrong answer overall -> NOR = 3 + 1 = 4."""
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 3)
        outcome = service.submit_attempt("wrong")
        assert outcome.remaining_repetitions == 4

    def test_requested_example_one_correct_two_wrong(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """NODR=3, one correct then two wrong -> NOR = 3 - 1 + 2 = 4."""
        seed(repository, "IELTS", ["alpha"])
        session = service.start_session(["IELTS"], 3)
        service.submit_attempt("alpha")  # correct; word returns to the pool
        service.submit_attempt("wrong")
        service.acknowledge_and_advance()
        outcome = service.submit_attempt("wrong")
        assert outcome.remaining_repetitions == 4
        assert session.pool_size == 1  # still not mastered, still in the pool

    def test_a_word_mastered_under_the_original_target_can_still_need_more(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """3 correct answers satisfy a required-3 session with a clean run; one
        mistake anywhere in the sequence means a 4th correct answer is needed."""
        seed(repository, "IELTS", ["alpha"])
        session = service.start_session(["IELTS"], 3)

        service.submit_attempt("alpha")  # correct_count=1
        service.submit_attempt("wrong")  # wrong_count=1
        service.acknowledge_and_advance()
        third = service.submit_attempt("alpha")  # correct_count=2
        # required(3) + wrong(1) - correct(2) = 2 remaining
        assert third.became_mastered is False
        assert third.remaining_repetitions == 2
        assert session.is_finished is False

        fourth = service.submit_attempt("alpha")  # correct_count=3, still 1 remaining
        assert fourth.became_mastered is False
        assert fourth.remaining_repetitions == 1
        assert session.is_finished is False

        fifth = service.submit_attempt("alpha")  # correct_count=4, 0 remaining
        assert fifth.became_mastered is True
        assert session.is_finished is True

    def test_word_stays_in_the_pool_after_a_miss(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        session = service.start_session(["IELTS"], 3)
        service.submit_attempt("wrong")
        assert session.pool_size == 1
        assert session.is_finished is False

    def test_drill_waits_for_acknowledgement_after_a_miss(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """FR-6.9: the user needs time to read the revealed spelling."""
        seed(repository, "IELTS", ["alpha", "beta"])
        session = service.start_session(["IELTS"], 3)

        service.submit_attempt("wrong")
        assert session.awaiting_acknowledgement is True

        service.acknowledge_and_advance()
        assert session.awaiting_acknowledgement is False

    def test_advancing_rotates_past_the_missed_word(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha", "beta", "gamma"])
        session = service.start_session(["IELTS"], 3)
        missed = session.current.entry.word  # type: ignore[union-attr]

        service.submit_attempt("wrong")
        service.acknowledge_and_advance()

        assert session.current.entry.word != missed  # type: ignore[union-attr]

    def test_single_word_pool_redraws_the_same_word(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """E6-S3 branch 3a: nothing else to show, so it must not stall."""
        seed(repository, "IELTS", ["lonely"])
        service.start_session(["IELTS"], 3)
        service.submit_attempt("wrong")
        assert service.acknowledge_and_advance().entry.word == "lonely"  # type: ignore[union-attr]

    def test_score_can_go_negative(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 5)
        for _ in range(3):
            service.submit_attempt("wrong")
            service.acknowledge_and_advance()
        assert service.session.score == -3  # type: ignore[union-attr]

    def test_blank_answer_is_judged_incorrect(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """The drill screen filters these out first; this is the backstop."""
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 3)
        assert service.submit_attempt("").was_correct is False


class TestProgress:
    def test_attempts_accumulate_including_misses(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 5)

        service.submit_attempt("alpha")
        service.submit_attempt("wrong")
        service.acknowledge_and_advance()

        assert service.session.attempts_made == 2  # type: ignore[union-attr]

    def test_denominator_does_not_shrink_as_words_are_mastered(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """E6-S5 branch 2a: the bar must only ever advance."""
        seed(repository, "IELTS", ["alpha", "beta"])
        session = service.start_session(["IELTS"], 1)
        before = session.total_required_attempts

        service.submit_attempt(session.current.entry.word)  # type: ignore[union-attr]

        assert session.total_required_attempts == before


class TestSessionEnd:
    def test_summary_reports_the_session(
        self,
        service: PracticeService,
        repository: MasterFileRepository,
        clock: FixedClock,
    ) -> None:
        seed(repository, "IELTS", ["alpha", "beta"])
        session = service.start_session(["IELTS"], 1)

        service.submit_attempt(session.current.entry.word)  # type: ignore[union-attr]
        service.submit_attempt("wrong")
        service.acknowledge_and_advance()
        clock.advance(minutes=18, seconds=42)

        summary = service.end_session()

        assert summary.score == 0  # 1 correct - 1 penalty
        assert summary.total_penalties == 1
        assert summary.words_practiced == 1
        assert summary.elapsed == timedelta(minutes=18, seconds=42)

    def test_summary_is_recorded_to_history(
        self,
        service: PracticeService,
        repository: MasterFileRepository,
        history: HistoryStore,
    ) -> None:
        """FR-6.16."""
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 1)
        service.submit_attempt("alpha")
        service.end_session()

        assert len(history.all_sessions()) == 1

    def test_nothing_is_written_to_the_workbook(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """FR-6.16: the workbook stays a clean vocabulary list."""
        seed(repository, "IELTS", ["alpha"])
        before = repository.all_words()
        service.start_session(["IELTS"], 1)
        service.submit_attempt("alpha")
        service.end_session()
        assert repository.all_words() == before

    def test_ending_early_reports_only_completed_work(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """E6-S6 branch 1a."""
        seed(repository, "IELTS", ["alpha", "beta", "gamma"])
        session = service.start_session(["IELTS"], 1)
        service.submit_attempt(session.current.entry.word)  # type: ignore[union-attr]

        summary = service.end_session()
        assert summary.words_practiced == 1

    def test_history_failure_does_not_block_the_summary(
        self,
        service: PracticeService,
        repository: MasterFileRepository,
        history: HistoryStore,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """E6-S6 branch 4a."""
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 1)
        service.submit_attempt("alpha")

        monkeypatch.setattr(history, "append", lambda _s: False)
        assert service.end_session().score == 1

    def test_ending_without_a_session_raises(self, service: PracticeService) -> None:
        with pytest.raises(EmptyPoolError):
            service.end_session()


class TestPracticeLog:
    """The separate Excel log of sessions (user-requested), distinct from both
    the master workbook (never touched, FR-6.16) and the JSON history store."""

    @pytest.fixture
    def practice_log(self, tmp_path) -> PracticeLogRepository:  # type: ignore[no-untyped-def]
        return PracticeLogRepository(tmp_path / "practice-log.xlsx")

    @pytest.fixture
    def logged_service(
        self,
        repository: MasterFileRepository,
        history: HistoryStore,
        practice_log: PracticeLogRepository,
        fake_tts: FakeTts,
        fake_audio: FakeAudio,
        clock: FixedClock,
    ) -> PracticeService:
        return PracticeService(
            repository,
            history,
            practice_log=practice_log,
            tts=fake_tts,
            audio=fake_audio,
            clock=clock,
            rng=random.Random(1234),
        )

    def _log_rows(self, practice_log: PracticeLogRepository) -> list[tuple]:
        from openpyxl import load_workbook

        workbook = load_workbook(practice_log.path)
        sheet = workbook["Practice Log"]
        return [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)]

    def test_a_completed_session_logs_finish_true(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        logged_service.start_session(["IELTS"], 1)
        logged_service.submit_attempt("alpha")
        logged_service.end_session()

        rows = self._log_rows(practice_log)
        assert len(rows) == 1
        assert rows[0][5] is True  # Finish column

    def test_ending_early_logs_finish_false(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
    ) -> None:
        """FR-6.18: the pool still held words when the session stopped."""
        seed(repository, "IELTS", ["alpha", "beta", "gamma"])
        session = logged_service.start_session(["IELTS"], 1)
        logged_service.submit_attempt(session.current.entry.word)  # type: ignore[union-attr]
        logged_service.end_session()

        rows = self._log_rows(practice_log)
        assert rows[0][5] is False

    def test_logs_the_chosen_collections(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        seed(repository, "Daily", ["beta"])
        logged_service.start_session(["IELTS", "Daily"], 1)
        logged_service.submit_attempt("alpha")
        logged_service.submit_attempt("beta")
        logged_service.end_session()

        rows = self._log_rows(practice_log)
        assert rows[0][3] == "IELTS, Daily"

    def test_logs_the_total_penalty_count(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        logged_service.start_session(["IELTS"], 3)
        logged_service.submit_attempt("wrong")
        logged_service.submit_attempt("wrong")
        logged_service.acknowledge_and_advance()
        logged_service.end_session()

        rows = self._log_rows(practice_log)
        assert rows[0][4] == 2

    def test_logs_per_word_miss_counts_in_the_requested_format(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
    ) -> None:
        seed(repository, "IELTS", ["hello", "goodbye"])
        session = logged_service.start_session(["IELTS"], 5)

        # Deliberately miss whichever word is current, tallying by name.
        misses_needed = {"hello": 2, "goodbye": 3}
        remaining = dict(misses_needed)
        guard = 0
        while any(remaining.values()) and guard < 100:
            guard += 1
            word = session.current.entry.word  # type: ignore[union-attr]
            if remaining.get(word, 0) > 0:
                logged_service.submit_attempt("wrong")
                remaining[word] -= 1
                logged_service.acknowledge_and_advance()
            else:
                logged_service.submit_attempt(word)

        logged_service.end_session()

        rows = self._log_rows(practice_log)
        cell = rows[0][6]
        parts = set(cell.split("|"))
        assert parts == {"hello_2", "goodbye_3"}

    def test_a_perfect_session_logs_no_words_with_penalty(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        logged_service.start_session(["IELTS"], 1)
        logged_service.submit_attempt("alpha")
        logged_service.end_session()

        rows = self._log_rows(practice_log)
        assert rows[0][6] in (None, "")

    def test_start_and_end_time_are_logged(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
        clock: FixedClock,
    ) -> None:
        from datetime import time

        seed(repository, "IELTS", ["alpha"])
        logged_service.start_session(["IELTS"], 1)
        clock.advance(minutes=18, seconds=42)
        logged_service.submit_attempt("alpha")
        logged_service.end_session()

        rows = self._log_rows(practice_log)
        assert rows[0][1] == time(10, 24, 0)  # FIXED_NOW from conftest
        assert rows[0][2] == time(10, 42, 42)

    def test_missing_practice_log_is_optional(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """The default ``service`` fixture has no practice_log wired -- ending a
        session must not raise just because this feature is not configured."""
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 1)
        service.submit_attempt("alpha")
        service.end_session()  # must not raise

    def test_a_failed_log_write_does_not_block_the_summary(
        self,
        logged_service: PracticeService,
        repository: MasterFileRepository,
        practice_log: PracticeLogRepository,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        logged_service.start_session(["IELTS"], 1)
        logged_service.submit_attempt("alpha")

        monkeypatch.setattr(practice_log, "append_entry", lambda _e: False)
        assert logged_service.end_session().score == 1


class TestWordsForCollectionPreview:
    """FR-5.10: the setup screen can preview a collection before committing it."""

    def test_returns_only_that_collections_words(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha", "beta"])
        seed(repository, "Other", ["gamma"])

        words = service.words_for_collection("IELTS")

        assert sorted(w.word for w in words) == ["alpha", "beta"]

    def test_matches_case_insensitively(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        assert len(service.words_for_collection("  ielts  ")) == 1

    def test_empty_collection_returns_an_empty_list(
        self, service: PracticeService
    ) -> None:
        assert service.words_for_collection("General") == []

    def test_unknown_collection_returns_empty_rather_than_raising(
        self, service: PracticeService
    ) -> None:
        assert service.words_for_collection("nonexistent") == []

    def test_agrees_with_the_settings_preview(
        self,
        service: PracticeService,
        repository: MasterFileRepository,
        preferences_store,  # type: ignore[no-untyped-def]
        fake_tts: FakeTts,
        clock: FixedClock,
    ) -> None:
        """Both previews must read through one definition of collection membership --
        that is why the filter lives in the repository, not in each service."""
        from vocabulary_trainer.services.settings_service import SettingsService
        from vocabulary_trainer.services.widget_state_service import (
            WidgetStateService,
        )

        seed(repository, "IELTS", ["alpha", "beta"])
        settings = SettingsService(
            repository,
            preferences_store,
            WidgetStateService(preferences_store),
            tts=fake_tts,
            clock=clock,
        )

        assert service.words_for_collection("IELTS") == settings.words_for_collection(
            "IELTS"
        )


class TestPracticeAgain:
    def test_reuses_the_same_collections_and_threshold(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """FR-6.15."""
        seed(repository, "IELTS", ["alpha", "beta"])
        service.start_session(["IELTS"], 4)
        service.end_session()

        again = service.practice_again()

        assert again.required_correct == 4
        assert again.collection_names == ("IELTS",)
        assert again.initial_pool_size == 2

    def test_last_config_is_reported(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 5)
        assert service.last_session_config() == (("IELTS",), 5)

    def test_practice_again_without_history_raises(
        self, service: PracticeService
    ) -> None:
        with pytest.raises(EmptyPoolError):
            service.practice_again()

    def test_deleted_collections_raise_rather_than_starting_empty(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """E6-S6 branch 6a."""
        seed(repository, "Doomed", ["alpha"])
        service.start_session(["Doomed"], 1)
        service.end_session()

        repository.delete_collection("Doomed")

        with pytest.raises(EmptyPoolError):
            service.practice_again()


class TestOptionalDependencies:
    def test_works_without_tts_or_audio(
        self, repository: MasterFileRepository, history: HistoryStore, clock: FixedClock
    ) -> None:
        """A machine with no audio device must still be able to practise."""
        service = PracticeService(repository, history, clock=clock, rng=random.Random(1))
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 1)
        assert service.submit_attempt("alpha").was_correct is True

    def test_unavailable_tts_does_not_interrupt(
        self, repository: MasterFileRepository, history: HistoryStore, clock: FixedClock
    ) -> None:
        service = PracticeService(
            repository,
            history,
            tts=FakeTts(available=False),
            audio=FakeAudio(),
            clock=clock,
            rng=random.Random(1),
        )
        seed(repository, "IELTS", ["alpha"])
        service.start_session(["IELTS"], 1)
        assert service.submit_attempt("alpha").was_correct is True


class TestFullSessionRun:
    def test_a_session_terminates_and_masters_everything(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        """End-to-end: every word must eventually leave the pool."""
        seed(repository, "IELTS", [f"word{i:02d}" for i in range(8)])
        session = service.start_session(["IELTS"], 3)

        guard = 0
        while not session.is_finished and guard < 500:
            guard += 1
            state = session.current
            assert state is not None
            service.submit_attempt(state.entry.word)

        assert session.is_finished is True
        assert session.words_mastered == 8
        assert session.score == 24  # 8 words x 3 correct writes, no penalties

    def test_a_session_with_misses_still_terminates(
        self, service: PracticeService, repository: MasterFileRepository
    ) -> None:
        seed(repository, "IELTS", [f"word{i}" for i in range(5)])
        session = service.start_session(["IELTS"], 2)

        guard = 0
        miss_next = True
        while not session.is_finished and guard < 500:
            guard += 1
            state = session.current
            assert state is not None
            if miss_next:
                service.submit_attempt("definitely wrong")
                service.acknowledge_and_advance()
            else:
                service.submit_attempt(state.entry.word)
            miss_next = not miss_next

        assert session.is_finished is True
        assert session.words_mastered == 5
        assert session.penalties > 0
