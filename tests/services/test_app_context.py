"""Tests for the composition root (U10).

Constructs the real object graph -- every service, store, and integration client --
with no Qt application object in the process. That the whole graph builds and runs
headless is the clearest evidence that the UI layer holds no business logic.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from vocabulary_trainer.data.master_file_repository import (
    DEFAULT_COLLECTION_NAME,
    WORDS_HEADERS,
    WORDS_SHEET,
)
from vocabulary_trainer.domain.models import PartOfSpeech
from vocabulary_trainer.services.app_context import AppContext
from vocabulary_trainer.services.clock import FixedClock

from conftest import FIXED_NOW


@pytest.fixture
def context(tmp_path: Path) -> AppContext:
    """A context on a throwaway data directory, with no file watcher thread."""
    ctx = AppContext(
        app_data_dir=tmp_path / "appdata",
        clock=FixedClock(FIXED_NOW),
        start_watching=False,
    )
    yield ctx
    ctx.shutdown()


class TestStartup:
    def test_builds_the_entire_graph(self, context: AppContext) -> None:
        assert context.repository is not None
        assert context.collect is not None
        assert context.practice is not None
        assert context.settings is not None
        assert context.widget_state is not None

    def test_creates_the_data_directory(self, context: AppContext) -> None:
        assert context.app_data_dir.is_dir()

    def test_creates_a_usable_workbook(self, context: AppContext) -> None:
        """FR-8.1: the app must be ready to use on first launch."""
        assert context.repository.path.exists()
        assert context.startup_report.created_workbook is True

    def test_seeds_the_default_collection(self, context: AppContext) -> None:
        assert DEFAULT_COLLECTION_NAME in context.collect.available_collections()

    def test_first_run_produces_no_repair_notice(self, context: AppContext) -> None:
        """A fresh workbook is expected, not something to warn about."""
        assert context.startup_report.needs_user_notice is False

    def test_applies_the_saved_tts_voice(self, tmp_path: Path) -> None:
        data_dir = tmp_path / "appdata"
        first = AppContext(app_data_dir=data_dir, start_watching=False)
        first.settings.set_tts_voice("voice-en-GB")
        first.shutdown()

        second = AppContext(app_data_dir=data_dir, start_watching=False)
        try:
            assert second.preferences.tts_voice_id == "voice-en-GB"
        finally:
            second.shutdown()

    def test_applies_the_saved_merriam_webster_key(self, tmp_path: Path) -> None:
        """The key must reach the client, or the source stays silently disabled."""
        data_dir = tmp_path / "appdata"
        first = AppContext(app_data_dir=data_dir, start_watching=False)
        first.settings.set_merriam_webster_key("abc-123")
        first.shutdown()

        second = AppContext(app_data_dir=data_dir, start_watching=False)
        try:
            assert second.merriam_webster.is_available() is True
        finally:
            second.shutdown()


class TestSharedState:
    def test_settings_and_widget_state_share_one_preferences_object(
        self, context: AppContext
    ) -> None:
        """Two copies would let the two services drift apart."""
        assert context.settings.preferences is context.widget_state.preferences
        assert context.settings.preferences is context.preferences

    def test_a_widget_change_is_visible_through_settings(
        self, context: AppContext
    ) -> None:
        context.widget_state.set_visible(False)
        assert context.settings.preferences.widget_visible is False


class TestPersistenceAcrossRestarts:
    def test_preferences_survive_a_restart(self, tmp_path: Path) -> None:
        data_dir = tmp_path / "appdata"

        first = AppContext(app_data_dir=data_dir, start_watching=False)
        first.widget_state.move_to(1234, 567)
        first.widget_state.set_visible(False)
        first.shutdown()

        second = AppContext(app_data_dir=data_dir, start_watching=False)
        try:
            assert second.widget_state.position == (1234, 567)
            assert second.widget_state.is_visible is False
        finally:
            second.shutdown()

    def test_collected_words_survive_a_restart(self, tmp_path: Path) -> None:
        data_dir = tmp_path / "appdata"

        first = AppContext(app_data_dir=data_dir, start_watching=False)
        first.collect.save_word(
            "ubiquitous", PartOfSpeech.ADJECTIVE, DEFAULT_COLLECTION_NAME
        )
        first.shutdown()

        second = AppContext(app_data_dir=data_dir, start_watching=False)
        try:
            assert [e.word for e in second.repository.all_words()] == ["ubiquitous"]
        finally:
            second.shutdown()


class TestReloadNotification:
    def test_listeners_are_called_on_an_external_change(
        self, context: AppContext
    ) -> None:
        calls: list[int] = []
        context.on_data_reloaded(lambda: calls.append(1))

        context._on_external_change()

        assert calls == [1]

    def test_unsubscribe_stops_notifications(self, context: AppContext) -> None:
        calls: list[int] = []
        unsubscribe = context.on_data_reloaded(lambda: calls.append(1))

        unsubscribe()
        context._on_external_change()

        assert calls == []

    def test_a_raising_listener_does_not_block_the_others(
        self, context: AppContext
    ) -> None:
        """One view failing to refresh must not end all future reloads."""
        reached: list[str] = []
        context.on_data_reloaded(lambda: (_ for _ in ()).throw(RuntimeError("boom")))
        context.on_data_reloaded(lambda: reached.append("second"))

        context._on_external_change()

        assert reached == ["second"]

    def test_reload_picks_up_a_hand_edited_workbook(
        self, context: AppContext
    ) -> None:
        """E8-S5: an edit made in Excel must become visible."""
        from openpyxl import load_workbook

        assert context.repository.all_words() == []

        workbook = load_workbook(context.repository.path)
        workbook[WORDS_SHEET].append(
            ["handwritten", "noun", "added in Excel", "", DEFAULT_COLLECTION_NAME, None]
        )
        workbook.save(context.repository.path)
        workbook.close()

        context._on_external_change()

        assert [e.word for e in context.repository.all_words()] == ["handwritten"]


class TestWorkbookRepairOnStartup:
    def test_reports_a_repaired_header(self, tmp_path: Path) -> None:
        """FR-8.4: the user is told exactly what changed."""
        import json

        from openpyxl import Workbook

        data_dir = tmp_path / "appdata"
        data_dir.mkdir(parents=True)
        broken = data_dir / "hand-edited.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = WORDS_SHEET
        # A user renamed the first column in Excel.
        sheet.append(["Word", *WORDS_HEADERS[1:]])
        workbook.save(broken)
        workbook.close()

        # Write preferences directly so the context loads the broken workbook on its
        # very first startup -- constructing a context to set the path would repair the
        # file before the assertion could observe it.
        (data_dir / "preferences.json").write_text(
            json.dumps({"master_file_path": str(broken)}), encoding="utf-8"
        )

        context = AppContext(app_data_dir=data_dir, start_watching=False)
        try:
            assert context.repository.path == broken
            assert context.startup_report.needs_user_notice is True
            assert any("Words" in change for change in context.startup_report.changes)
        finally:
            context.shutdown()

    def test_a_user_chosen_path_is_honoured_over_the_default(
        self, tmp_path: Path
    ) -> None:
        """A location picked in Settings must survive a restart."""
        import json

        data_dir = tmp_path / "appdata"
        data_dir.mkdir(parents=True)
        elsewhere = tmp_path / "my-vocab" / "words.xlsx"
        (data_dir / "preferences.json").write_text(
            json.dumps({"master_file_path": str(elsewhere)}), encoding="utf-8"
        )

        context = AppContext(app_data_dir=data_dir, start_watching=False)
        try:
            assert context.repository.path == elsewhere
            assert elsewhere.exists()
        finally:
            context.shutdown()


class TestShutdown:
    def test_is_idempotent(self, tmp_path: Path) -> None:
        context = AppContext(app_data_dir=tmp_path / "appdata", start_watching=False)
        context.shutdown()
        context.shutdown()

    def test_persists_preferences(self, tmp_path: Path) -> None:
        data_dir = tmp_path / "appdata"
        context = AppContext(app_data_dir=data_dir, start_watching=False)
        context.preferences.required_correct_writes = 7
        context.shutdown()

        reopened = AppContext(app_data_dir=data_dir, start_watching=False)
        try:
            assert reopened.preferences.required_correct_writes == 7
        finally:
            reopened.shutdown()

    def test_leaves_no_temp_files_behind(self, tmp_path: Path) -> None:
        data_dir = tmp_path / "appdata"
        context = AppContext(app_data_dir=data_dir, start_watching=False)
        context.collect.save_word("word", PartOfSpeech.NOUN, DEFAULT_COLLECTION_NAME)
        context.shutdown()

        leftovers = [p.name for p in data_dir.iterdir() if ".tmp" in p.name]
        assert leftovers == []


class TestEndToEndFlows:
    def test_collect_then_practice(self, context: AppContext) -> None:
        """The two pillars of the product, exercised through the real graph."""
        for word in ("alpha", "beta", "gamma"):
            outcome = context.collect.save_word(
                word, PartOfSpeech.NOUN, DEFAULT_COLLECTION_NAME
            )
            assert outcome.status.value == "saved"

        assert context.practice.can_start([DEFAULT_COLLECTION_NAME]) is True
        session = context.practice.start_session([DEFAULT_COLLECTION_NAME], 1)
        assert session.initial_pool_size == 3

        guard = 0
        while not session.is_finished and guard < 100:
            guard += 1
            state = session.current
            assert state is not None
            context.practice.submit_attempt(state.entry.word)

        summary = context.practice.end_session()
        assert summary.words_practiced == 3
        assert summary.score == 3

    def test_collect_then_rename_the_collection(self, context: AppContext) -> None:
        context.collect.create_collection_inline("IELTS")
        context.collect.save_word("ubiquitous", PartOfSpeech.ADJECTIVE, "IELTS")

        context.settings.rename_collection("IELTS", "IELTS Advanced")

        assert context.repository.all_words()[0].collection_name == "IELTS Advanced"

    def test_collect_then_delete_the_collection(self, context: AppContext) -> None:
        context.collect.create_collection_inline("Doomed")
        context.collect.save_word("victim", PartOfSpeech.NOUN, "Doomed")

        promised = context.settings.word_count_for_deletion("Doomed")
        removed = context.settings.delete_collection("Doomed")

        assert removed == promised == 1
        assert context.repository.all_words() == []
