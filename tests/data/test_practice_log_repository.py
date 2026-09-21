"""Tests for the separate practice-log workbook (user-requested feature).

Distinct from ``HistoryStore`` (JSON, capped, score-focused) and from the master
workbook (never touched by practice at all, per FR-6.16) -- this is a third store,
purely additive, aimed at a spreadsheet-native record of session timing, which
collections were drilled, and which specific words needed correction.
"""

from __future__ import annotations

from datetime import date, time
from pathlib import Path

from openpyxl import load_workbook

from vocabulary_trainer.data.practice_log_repository import (
    HEADERS,
    SHEET_NAME,
    PracticeLogRepository,
)
from vocabulary_trainer.domain.models import PracticeLogEntry


def entry(
    when: date = date(2026, 9, 14),
    start: time = time(10, 0, 0),
    end: time = time(10, 18, 42),
    collections: tuple[str, ...] = ("General",),
    penalty: int = 0,
    finished: bool = True,
    wrong_counts: tuple[tuple[str, int], ...] = (),
) -> PracticeLogEntry:
    return PracticeLogEntry(
        session_date=when,
        start_time=start,
        end_time=end,
        collections=collections,
        penalty=penalty,
        finished=finished,
        wrong_counts=wrong_counts,
    )


class TestAppendEntry:
    def test_creates_the_file_on_first_append(self, tmp_path: Path) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        assert repo.append_entry(entry()) is True
        assert repo.path.is_file()

    def test_writes_the_expected_header_row(self, tmp_path: Path) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        repo.append_entry(entry())

        workbook = load_workbook(repo.path)
        sheet = workbook[SHEET_NAME]
        assert tuple(cell.value for cell in sheet[1]) == HEADERS

    def test_row_content_matches_the_entry(self, tmp_path: Path) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        repo.append_entry(
            entry(
                when=date(2026, 9, 14),
                start=time(10, 0, 0),
                end=time(10, 18, 42),
                collections=("IELTS", "Daily Reading"),
                penalty=5,
                finished=False,
                wrong_counts=(("goodbye", 3), ("hello", 2)),
            )
        )

        workbook = load_workbook(repo.path)
        sheet = workbook[SHEET_NAME]
        row = tuple(cell.value for cell in sheet[2])
        assert row[0].date() == date(2026, 9, 14)
        assert row[1] == time(10, 0, 0)
        assert row[2] == time(10, 18, 42)
        assert row[3] == "IELTS, Daily Reading"
        assert row[4] == 5
        assert row[5] is False
        assert row[6] == "goodbye_3|hello_2"

    def test_a_perfect_session_leaves_the_penalty_word_cell_blank(
        self, tmp_path: Path
    ) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        repo.append_entry(entry(penalty=0, wrong_counts=()))

        workbook = load_workbook(repo.path)
        sheet = workbook[SHEET_NAME]
        assert sheet.cell(row=2, column=7).value in (None, "")

    def test_finished_session_is_true(self, tmp_path: Path) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        repo.append_entry(entry(finished=True))
        workbook = load_workbook(repo.path)
        assert workbook[SHEET_NAME].cell(row=2, column=6).value is True

    def test_time_columns_use_an_hms_number_format(self, tmp_path: Path) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        repo.append_entry(entry())
        workbook = load_workbook(repo.path)
        sheet = workbook[SHEET_NAME]
        assert sheet.cell(row=2, column=2).number_format == "hh:mm:ss"
        assert sheet.cell(row=2, column=3).number_format == "hh:mm:ss"

    def test_appends_without_overwriting_earlier_rows(self, tmp_path: Path) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        repo.append_entry(entry(collections=("First",)))
        repo.append_entry(entry(collections=("Second",)))

        workbook = load_workbook(repo.path)
        sheet = workbook[SHEET_NAME]
        assert sheet.cell(row=2, column=4).value == "First"
        assert sheet.cell(row=3, column=4).value == "Second"

    def test_a_single_collection_has_no_stray_separator(
        self, tmp_path: Path
    ) -> None:
        repo = PracticeLogRepository(tmp_path / "practice-log.xlsx")
        repo.append_entry(entry(collections=("Solo",)))
        workbook = load_workbook(repo.path)
        assert workbook[SHEET_NAME].cell(row=2, column=4).value == "Solo"

    def test_recovers_from_a_corrupt_existing_file_rather_than_failing(
        self, tmp_path: Path
    ) -> None:
        path = tmp_path / "practice-log.xlsx"
        path.write_bytes(b"not a real xlsx file")
        repo = PracticeLogRepository(path)

        assert repo.append_entry(entry()) is True
        workbook = load_workbook(path)
        assert workbook[SHEET_NAME].max_row == 2  # header + the new row


class TestPracticeLogEntryFormatting:
    def test_word_with_penalty_matches_the_requested_format(self) -> None:
        e = entry(wrong_counts=(("hello", 2), ("goodbye", 3)))
        assert e.word_with_penalty == "hello_2|goodbye_3"

    def test_no_misses_is_an_empty_string(self) -> None:
        assert entry(wrong_counts=()).word_with_penalty == ""

    def test_single_miss_has_no_separator(self) -> None:
        assert entry(wrong_counts=(("hello", 1),)).word_with_penalty == "hello_1"

    def test_collection_display_joins_with_a_comma(self) -> None:
        assert (
            entry(collections=("A", "B", "C")).collection_display == "A, B, C"
        )

    def test_single_collection_display_has_no_comma(self) -> None:
        assert entry(collections=("Solo",)).collection_display == "Solo"
