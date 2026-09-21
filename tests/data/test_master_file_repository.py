"""Tests for the Excel master file repository.

These run against real temporary ``.xlsx`` files rather than mocks. This is the
unit where data loss would actually happen, so the tests exercise the genuine
openpyxl round-trip, the atomic swap, and the cascade operations (NFR-TEST-03).
"""

from __future__ import annotations

import threading
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from vocabulary_trainer.data.errors import (
    CollectionNotFoundError,
    DuplicateCollectionError,
    MasterFileLockedError,
    MasterFileUnreadableError,
)
from vocabulary_trainer.data.master_file_repository import (
    COLLECTIONS_HEADERS,
    COLLECTIONS_SHEET,
    DEFAULT_COLLECTION_NAME,
    WORDS_HEADERS,
    WORDS_SHEET,
    MasterFileRepository,
)
from vocabulary_trainer.domain.models import Collection, PartOfSpeech, WordEntry


@pytest.fixture
def repo_path(tmp_path: Path) -> Path:
    return tmp_path / "master.xlsx"


@pytest.fixture
def repo(repo_path: Path) -> MasterFileRepository:
    repository = MasterFileRepository(repo_path)
    repository.ensure_workbook()
    return repository


def entry(
    word: str = "ubiquitous",
    collection: str = DEFAULT_COLLECTION_NAME,
    meaning: str | None = "present everywhere",
    example: str | None = "Phones are ubiquitous.",
    pos: PartOfSpeech = PartOfSpeech.ADJECTIVE,
    when: date | None = None,
) -> WordEntry:
    return WordEntry(
        word=word,
        part_of_speech=pos,
        meaning=meaning,
        example=example,
        collection_name=collection,
        collected_on=when or date(2026, 8, 14),
    )


class TestFirstRunCreation:
    def test_creates_workbook_when_absent(self, repo_path: Path) -> None:
        report = MasterFileRepository(repo_path).ensure_workbook()
        assert repo_path.exists()
        assert report.created_workbook is True

    def test_creates_parent_directories(self, tmp_path: Path) -> None:
        nested = tmp_path / "a" / "b" / "c" / "master.xlsx"
        MasterFileRepository(nested).ensure_workbook()
        assert nested.exists()

    def test_creates_both_sheets_with_correct_headers(self, repo_path: Path) -> None:
        MasterFileRepository(repo_path).ensure_workbook()
        workbook = load_workbook(repo_path)
        assert WORDS_SHEET in workbook.sheetnames
        assert COLLECTIONS_SHEET in workbook.sheetnames
        assert tuple(c.value for c in workbook[WORDS_SHEET][1]) == WORDS_HEADERS
        assert tuple(c.value for c in workbook[COLLECTIONS_SHEET][1]) == COLLECTIONS_HEADERS

    def test_seeds_the_default_collection(self, repo: MasterFileRepository) -> None:
        """FR-8.1: the very first Collect action must have a destination."""
        names = [c.name for c in repo.all_collections()]
        assert names == [DEFAULT_COLLECTION_NAME]

    def test_creation_is_not_reported_as_a_repair(self, repo_path: Path) -> None:
        """A fresh workbook is expected, not something to warn the user about."""
        report = MasterFileRepository(repo_path).ensure_workbook()
        assert report.needs_user_notice is False
        assert report.changes == ()

    def test_existing_workbook_is_not_overwritten(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry("preexisting"))
        reopened = MasterFileRepository(repo.path)
        reopened.ensure_workbook()
        assert [e.word for e in reopened.all_words()] == ["preexisting"]

    def test_seeds_general_into_a_workbook_that_has_none(self, repo_path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = WORDS_SHEET
        sheet.append(list(WORDS_HEADERS))
        sheet.append(["orphan", "noun", "m", "e", "Gone", date(2026, 1, 1)])
        collections = workbook.create_sheet(COLLECTIONS_SHEET)
        collections.append(list(COLLECTIONS_HEADERS))
        workbook.save(repo_path)

        repository = MasterFileRepository(repo_path)
        report = repository.ensure_workbook()

        assert report.seeded_default_collection is True
        assert [c.name for c in repository.all_collections()] == [DEFAULT_COLLECTION_NAME]
        # The orphaned word keeps its own collection name; seeding only guarantees
        # that Collect has a destination.
        assert repository.all_words()[0].collection_name == "Gone"


class TestHeaderRepair:
    def _workbook_with_words_header(self, path: Path, headers: list[object]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = WORDS_SHEET
        sheet.append(headers)
        sheet.append(["ubiquitous", "adjective", "m", "e", "General", date(2026, 5, 1)])
        collections = workbook.create_sheet(COLLECTIONS_SHEET)
        collections.append(list(COLLECTIONS_HEADERS))
        collections.append([DEFAULT_COLLECTION_NAME, date(2026, 1, 1)])
        workbook.save(path)

    def test_renamed_column_is_corrected_and_reported(self, repo_path: Path) -> None:
        self._workbook_with_words_header(
            repo_path,
            ["Word", "Type", "Meaning", "Example", "Name of collection", "Date"],
        )
        report = MasterFileRepository(repo_path).ensure_workbook()
        assert report.needs_user_notice is True
        assert any("renamed column 'Word'" in c for c in report.changes)
        workbook = load_workbook(repo_path)
        assert workbook[WORDS_SHEET]["A1"].value == "Words"

    def test_repair_preserves_the_data_under_the_column(self, repo_path: Path) -> None:
        """Repair is positional, so the row beneath a renamed header survives."""
        self._workbook_with_words_header(
            repo_path,
            ["Word", "Type", "Meaning", "Example", "Name of collection", "Date"],
        )
        repository = MasterFileRepository(repo_path)
        repository.ensure_workbook()
        words = repository.all_words()
        assert len(words) == 1
        assert words[0].word == "ubiquitous"

    def test_blank_column_is_added(self, repo_path: Path) -> None:
        self._workbook_with_words_header(
            repo_path, ["Words", "Type", None, "Example", "Name of collection", "Date"]
        )
        report = MasterFileRepository(repo_path).ensure_workbook()
        assert any("added the missing column 'Meaning'" in c for c in report.changes)

    def test_missing_header_row_is_added(self, repo_path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = WORDS_SHEET
        workbook.create_sheet(COLLECTIONS_SHEET).append(list(COLLECTIONS_HEADERS))
        workbook.save(repo_path)

        report = MasterFileRepository(repo_path).ensure_workbook()
        assert any("added the missing header row" in c for c in report.changes)

    def test_missing_sheet_is_created(self, repo_path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = WORDS_SHEET
        sheet.append(list(WORDS_HEADERS))
        workbook.save(repo_path)

        report = MasterFileRepository(repo_path).ensure_workbook()
        assert any(COLLECTIONS_SHEET in c for c in report.changes)
        assert COLLECTIONS_SHEET in load_workbook(repo_path).sheetnames

    def test_valid_headers_produce_no_changes(self, repo_path: Path) -> None:
        self._workbook_with_words_header(repo_path, list(WORDS_HEADERS))
        report = MasterFileRepository(repo_path).ensure_workbook()
        assert report.changes == ()
        assert report.needs_user_notice is False

    @pytest.mark.parametrize(
        "variant",
        [
            ["words", "type", "meaning", "example", "name of collection", "date"],
            [" Words ", "Type", "Meaning", "Example", "Name of collection", "Date"],
            ["WORDS", "TYPE", "MEANING", "EXAMPLE", "NAME OF COLLECTION", "DATE"],
        ],
    )
    def test_case_and_whitespace_variants_are_accepted_silently(
        self, repo_path: Path, variant: list[object]
    ) -> None:
        """Noisily 'repairing' a harmless variant would train users to ignore warnings."""
        self._workbook_with_words_header(repo_path, variant)
        report = MasterFileRepository(repo_path).ensure_workbook()
        assert report.changes == ()

    def test_extra_user_columns_are_preserved(self, repo_path: Path) -> None:
        """The workbook is hand-editable; destroying a user's own column would be
        indefensible."""
        self._workbook_with_words_header(
            repo_path, [*WORDS_HEADERS, "My personal notes"]
        )
        MasterFileRepository(repo_path).ensure_workbook()
        workbook = load_workbook(repo_path)
        assert workbook[WORDS_SHEET].cell(row=1, column=7).value == "My personal notes"


class TestUnreadableFile:
    def test_non_workbook_raises_and_is_not_repaired(self, tmp_path: Path) -> None:
        bogus = tmp_path / "master.xlsx"
        bogus.write_text("this is not a workbook", encoding="utf-8")
        with pytest.raises(MasterFileUnreadableError):
            MasterFileRepository(bogus).ensure_workbook()


class TestReads:
    def test_round_trips_a_word(self, repo: MasterFileRepository) -> None:
        original = entry()
        repo.insert_word(original)
        stored = repo.all_words()[0]
        assert stored.word == original.word
        assert stored.part_of_speech is original.part_of_speech
        assert stored.meaning == original.meaning
        assert stored.example == original.example
        assert stored.collection_name == original.collection_name
        assert stored.collected_on == original.collected_on

    def test_blank_meaning_and_example_read_back_as_none(
        self, repo: MasterFileRepository
    ) -> None:
        """A word saved mid-lookup has neither (FR-2.8)."""
        repo.insert_word(entry(meaning=None, example=None))
        stored = repo.all_words()[0]
        assert stored.meaning is None
        assert stored.example is None

    def test_collection_summaries_carry_counts(self, repo: MasterFileRepository) -> None:
        repo.insert_collection(Collection("IELTS", date(2026, 8, 14)))
        repo.insert_word(entry("alpha", "IELTS"))
        repo.insert_word(entry("beta", "IELTS"))
        repo.insert_word(entry("gamma", DEFAULT_COLLECTION_NAME))

        summaries = {s.name: s.word_count for s in repo.collection_summaries()}
        assert summaries == {"IELTS": 2, DEFAULT_COLLECTION_NAME: 1}

    def test_counts_match_case_insensitively(self, repo: MasterFileRepository) -> None:
        repo.insert_collection(Collection("IELTS", date(2026, 8, 14)))
        repo.insert_word(entry("alpha", "ielts"))
        summaries = {s.name: s.word_count for s in repo.collection_summaries()}
        assert summaries["IELTS"] == 1

    def test_word_count_for(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry("alpha"))
        repo.insert_word(entry("beta"))
        assert repo.word_count_for(DEFAULT_COLLECTION_NAME) == 2
        assert repo.word_count_for("nonexistent") == 0


class TestCellCoercion:
    """The workbook is hand-edited, so a formatting mistake must not cost a word."""

    def _raw_row(self, repo: MasterFileRepository, row: list[object]) -> None:
        workbook = load_workbook(repo.path)
        workbook[WORDS_SHEET].append(row)
        workbook.save(repo.path)
        workbook.close()
        repo.invalidate_cache()

    def test_blank_word_row_is_skipped(self, repo: MasterFileRepository) -> None:
        """The only condition that discards a row: it carries no information."""
        self._raw_row(repo, ["", "noun", "m", "e", "General", date(2026, 1, 1)])
        assert repo.all_words() == []

    def test_unknown_type_defaults_to_noun(self, repo: MasterFileRepository) -> None:
        self._raw_row(repo, ["word", "gibberish", "m", "e", "General", date(2026, 1, 1)])
        assert repo.all_words()[0].part_of_speech is PartOfSpeech.NOUN

    def test_blank_collection_becomes_general(self, repo: MasterFileRepository) -> None:
        self._raw_row(repo, ["word", "noun", "m", "e", "", date(2026, 1, 1)])
        assert repo.all_words()[0].collection_name == DEFAULT_COLLECTION_NAME

    @pytest.mark.parametrize(
        ("cell", "expected"),
        [
            ("2026-05-01", date(2026, 5, 1)),
            ("01/05/2026", date(2026, 5, 1)),
            (datetime(2026, 5, 1, 10, 30), date(2026, 5, 1)),
        ],
    )
    def test_date_formats_are_parsed(
        self, repo: MasterFileRepository, cell: object, expected: date
    ) -> None:
        self._raw_row(repo, ["word", "noun", "m", "e", "General", cell])
        assert repo.all_words()[0].collected_on == expected

    def test_unparseable_date_falls_back_to_today(
        self, repo: MasterFileRepository
    ) -> None:
        """A word with a broken date is still a word the user collected."""
        self._raw_row(repo, ["word", "noun", "m", "e", "General", "not a date"])
        assert repo.all_words()[0].collected_on == date.today()

    def test_type_matching_is_case_insensitive(self, repo: MasterFileRepository) -> None:
        self._raw_row(repo, ["word", "ADJECTIVE", "m", "e", "General", date(2026, 1, 1)])
        assert repo.all_words()[0].part_of_speech is PartOfSpeech.ADJECTIVE


class TestNaturalKeyLookup:
    def test_exact_match(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry("ubiquitous", "IELTS"))
        assert repo.find_by_natural_key("ubiquitous", "IELTS") is not None

    def test_case_difference_matches(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry("ubiquitous", "IELTS"))
        assert repo.find_by_natural_key("UBIQUITOUS", "ielts") is not None

    def test_whitespace_difference_matches(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry("give up", "IELTS"))
        assert repo.find_by_natural_key("  give   up  ", " IELTS ") is not None

    def test_no_match_returns_none(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry("ubiquitous", "IELTS"))
        assert repo.find_by_natural_key("different", "IELTS") is None

    def test_same_word_other_collection_is_not_a_match(
        self, repo: MasterFileRepository
    ) -> None:
        """The key is (word, collection) -- the same word twice is legitimate."""
        repo.insert_word(entry("ubiquitous", "IELTS"))
        assert repo.find_by_natural_key("ubiquitous", "Daily Reading") is None


class TestWordMutations:
    def test_patch_enrichment_fills_meaning_and_example(
        self, repo: MasterFileRepository
    ) -> None:
        repo.insert_word(entry(meaning=None, example=None))
        patched = repo.patch_word_enrichment(
            ("ubiquitous", DEFAULT_COLLECTION_NAME.casefold()),
            "present everywhere",
            "An example.",
        )
        assert patched is True
        stored = repo.all_words()[0]
        assert stored.meaning == "present everywhere"
        assert stored.example == "An example."

    def test_patch_of_a_deleted_row_returns_false(
        self, repo: MasterFileRepository
    ) -> None:
        """A lookup finishing after its row was removed must not resurrect it
        (E2-S4 branch 6c)."""
        patched = repo.patch_word_enrichment(("gone", "general"), "m", "e")
        assert patched is False
        assert repo.all_words() == []

    def test_replace_preserves_word_and_collection(
        self, repo: MasterFileRepository
    ) -> None:
        """FR-4.4: Replace refreshes content, it does not move or rename."""
        repo.insert_word(entry(meaning="old", example="old ex", when=date(2026, 1, 1)))
        repo.replace_word(
            entry(meaning="new", example="new ex", when=date(2026, 9, 12))
        )
        stored = repo.all_words()[0]
        assert stored.word == "ubiquitous"
        assert stored.collection_name == DEFAULT_COLLECTION_NAME
        assert stored.meaning == "new"
        assert stored.example == "new ex"
        assert stored.collected_on == date(2026, 9, 12)

    def test_replace_does_not_add_a_row(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry())
        repo.replace_word(entry(meaning="new"))
        assert len(repo.all_words()) == 1

    def test_delete_word(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry("alpha"))
        repo.insert_word(entry("beta"))
        deleted = repo.delete_word(("alpha", DEFAULT_COLLECTION_NAME.casefold()))
        assert deleted is True
        assert [e.word for e in repo.all_words()] == ["beta"]

    def test_delete_missing_word_returns_false(self, repo: MasterFileRepository) -> None:
        assert repo.delete_word(("nope", "general")) is False


class TestCascadeRename:
    def test_updates_every_belonging_word_row(self, repo: MasterFileRepository) -> None:
        repo.insert_collection(Collection("IELTS", date(2026, 8, 14)))
        for word in ("alpha", "beta", "gamma"):
            repo.insert_word(entry(word, "IELTS"))

        repo.rename_collection("IELTS", "IELTS Advanced")

        assert all(e.collection_name == "IELTS Advanced" for e in repo.all_words())
        assert "IELTS Advanced" in [c.name for c in repo.all_collections()]

    def test_leaves_other_collections_rows_untouched(
        self, repo: MasterFileRepository
    ) -> None:
        repo.insert_collection(Collection("IELTS", date(2026, 8, 14)))
        repo.insert_word(entry("alpha", "IELTS"))
        repo.insert_word(entry("beta", DEFAULT_COLLECTION_NAME))

        repo.rename_collection("IELTS", "Renamed")

        by_word = {e.word: e.collection_name for e in repo.all_words()}
        assert by_word == {"alpha": "Renamed", "beta": DEFAULT_COLLECTION_NAME}

    def test_recasing_own_name_is_permitted(self, repo: MasterFileRepository) -> None:
        """E7-S7 branch 4c: recasing is not a collision with another collection."""
        repo.insert_collection(Collection("ielts", date(2026, 8, 14)))
        repo.insert_word(entry("alpha", "ielts"))

        repo.rename_collection("ielts", "IELTS")

        assert "IELTS" in [c.name for c in repo.all_collections()]
        assert repo.all_words()[0].collection_name == "IELTS"

    def test_collision_with_another_collection_is_rejected(
        self, repo: MasterFileRepository
    ) -> None:
        repo.insert_collection(Collection("IELTS", date(2026, 8, 14)))
        repo.insert_collection(Collection("TOEFL", date(2026, 8, 15)))
        with pytest.raises(DuplicateCollectionError):
            repo.rename_collection("IELTS", "toefl")

    def test_missing_collection_raises(self, repo: MasterFileRepository) -> None:
        with pytest.raises(CollectionNotFoundError):
            repo.rename_collection("nonexistent", "whatever")

    def test_matching_is_case_insensitive_but_writing_is_verbatim(
        self, repo: MasterFileRepository
    ) -> None:
        repo.insert_collection(Collection("IELTS", date(2026, 8, 14)))
        repo.insert_word(entry("alpha", "ielts"))
        repo.rename_collection("IELTS", "My New Name")
        assert repo.all_words()[0].collection_name == "My New Name"


class TestCascadeDelete:
    def test_removes_the_collection_and_its_words(
        self, repo: MasterFileRepository
    ) -> None:
        repo.insert_collection(Collection("Business English", date(2026, 7, 20)))
        for word in ("alpha", "beta", "gamma"):
            repo.insert_word(entry(word, "Business English"))

        removed = repo.delete_collection("Business English")

        assert removed == 3
        assert repo.all_words() == []
        assert "Business English" not in [c.name for c in repo.all_collections()]

    def test_returns_the_exact_count_the_dialog_promised(
        self, repo: MasterFileRepository
    ) -> None:
        """FR-7.10 states an exact number; FR-7.11 must match it."""
        repo.insert_collection(Collection("Doomed", date(2026, 1, 1)))
        for index in range(17):
            repo.insert_word(entry(f"word{index}", "Doomed"))
        promised = repo.word_count_for("Doomed")
        assert repo.delete_collection("Doomed") == promised == 17

    def test_other_collections_are_unaffected(self, repo: MasterFileRepository) -> None:
        repo.insert_collection(Collection("Keep", date(2026, 1, 1)))
        repo.insert_collection(Collection("Doomed", date(2026, 1, 2)))
        repo.insert_word(entry("survivor", "Keep"))
        repo.insert_word(entry("victim", "Doomed"))

        repo.delete_collection("Doomed")

        assert [e.word for e in repo.all_words()] == ["survivor"]
        assert [c.name for c in repo.all_collections()] == [
            DEFAULT_COLLECTION_NAME,
            "Keep",
        ]

    def test_empty_collection_is_deletable(self, repo: MasterFileRepository) -> None:
        repo.insert_collection(Collection("Empty", date(2026, 1, 1)))
        assert repo.delete_collection("Empty") == 0

    def test_missing_collection_raises(self, repo: MasterFileRepository) -> None:
        with pytest.raises(CollectionNotFoundError):
            repo.delete_collection("nonexistent")


class TestCollectionInsertion:
    def test_duplicate_name_is_rejected_case_insensitively(
        self, repo: MasterFileRepository
    ) -> None:
        repo.insert_collection(Collection("IELTS", date(2026, 8, 14)))
        with pytest.raises(DuplicateCollectionError):
            repo.insert_collection(Collection("ielts", date(2026, 8, 15)))

    def test_ensure_default_is_a_noop_when_collections_exist(
        self, repo: MasterFileRepository
    ) -> None:
        assert repo.ensure_default_collection() is False

    def test_ensure_default_reseeds_after_the_last_collection_is_deleted(
        self, repo: MasterFileRepository
    ) -> None:
        """E7-S8 branch 7b: Collect must always have a destination."""
        repo.delete_collection(DEFAULT_COLLECTION_NAME)
        assert repo.all_collections() == []
        assert repo.ensure_default_collection() is True
        assert [c.name for c in repo.all_collections()] == [DEFAULT_COLLECTION_NAME]


class TestAtomicWrite:
    def test_no_temp_file_remains_after_success(self, repo: MasterFileRepository) -> None:
        repo.insert_word(entry())
        leftovers = list(repo.path.parent.glob(f".{repo.path.name}.tmp-*"))
        assert leftovers == []

    def test_failure_before_the_swap_leaves_the_original_intact(
        self, repo: MasterFileRepository, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        repo.insert_word(entry("original"))
        before = repo.path.read_bytes()

        def explode(*_args: object, **_kwargs: object) -> None:
            raise RuntimeError("simulated serialization failure")

        monkeypatch.setattr(Workbook, "save", explode)

        with pytest.raises(RuntimeError):
            repo.insert_word(entry("never lands"))

        assert repo.path.read_bytes() == before

    def test_failure_removes_the_temp_file(
        self, repo: MasterFileRepository, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def explode(*_args: object, **_kwargs: object) -> None:
            raise RuntimeError("simulated failure")

        monkeypatch.setattr(Workbook, "save", explode)
        with pytest.raises(RuntimeError):
            repo.insert_word(entry())

        assert list(repo.path.parent.glob(f".{repo.path.name}.tmp-*")) == []

    def test_orphaned_temp_files_are_cleaned_on_ensure(
        self, repo: MasterFileRepository
    ) -> None:
        """A process killed between save and swap leaves one behind."""
        orphan = repo.path.parent / f".{repo.path.name}.tmp-deadbeef"
        orphan.write_bytes(b"junk")
        repo.ensure_workbook()
        assert not orphan.exists()

    def test_temp_file_is_created_beside_the_target(
        self, repo: MasterFileRepository, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """os.replace is only atomic within one volume, so the temp file must be
        a sibling rather than living under %TEMP%."""
        seen: list[Path] = []
        original = Workbook.save

        def spy(self: Workbook, filename: object) -> None:  # type: ignore[override]
            seen.append(Path(str(filename)))
            original(self, filename)

        monkeypatch.setattr(Workbook, "save", spy)
        repo.insert_word(entry())

        assert seen
        assert all(p.parent == repo.path.parent for p in seen)


class TestLockedFile:
    def test_permission_error_becomes_master_file_locked(
        self, repo: MasterFileRepository, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """FR-8.6: the caller needs a distinguishable error to offer Retry."""

        def deny(*_args: object, **_kwargs: object) -> None:
            raise PermissionError(13, "in use by another process")

        monkeypatch.setattr("vocabulary_trainer.data.master_file_repository.os.replace", deny)

        with pytest.raises(MasterFileLockedError):
            repo.insert_word(entry())

    def test_locked_write_leaves_previous_content(
        self, repo: MasterFileRepository, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        repo.insert_word(entry("original"))
        before = repo.path.read_bytes()

        def deny(*_args: object, **_kwargs: object) -> None:
            raise PermissionError(13, "in use")

        monkeypatch.setattr("vocabulary_trainer.data.master_file_repository.os.replace", deny)

        with pytest.raises(MasterFileLockedError):
            repo.insert_word(entry("blocked"))

        assert repo.path.read_bytes() == before

    def test_locked_error_message_names_the_file_and_suggests_a_fix(
        self, repo: MasterFileRepository
    ) -> None:
        error = MasterFileLockedError(repo.path)
        assert str(repo.path) in str(error)
        assert "Excel" in str(error)


class TestWriteSerialization:
    def test_concurrent_inserts_all_land(self, repo: MasterFileRepository) -> None:
        """FR-8.7: without the lock, concurrent read-modify-write would lose rows."""
        count = 12
        errors: list[BaseException] = []

        def insert(index: int) -> None:
            try:
                repo.insert_word(entry(f"word{index:02d}"))
            except BaseException as exc:  # pragma: no cover - surfaced via assert
                errors.append(exc)

        threads = [threading.Thread(target=insert, args=(i,)) for i in range(count)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()

        assert errors == []
        assert len(repo.all_words()) == count

    def test_concurrent_mixed_operations_stay_consistent(
        self, repo: MasterFileRepository
    ) -> None:
        """E8-S6: overlapping activity must not scramble the file."""
        for index in range(6):
            repo.insert_word(entry(f"seed{index}"))

        errors: list[BaseException] = []

        def insert_more(index: int) -> None:
            try:
                repo.insert_word(entry(f"extra{index}"))
            except BaseException as exc:  # pragma: no cover
                errors.append(exc)

        def add_collection(index: int) -> None:
            try:
                repo.insert_collection(Collection(f"C{index}", date(2026, 1, 1)))
            except BaseException as exc:  # pragma: no cover
                errors.append(exc)

        threads = [
            *(threading.Thread(target=insert_more, args=(i,)) for i in range(4)),
            *(threading.Thread(target=add_collection, args=(i,)) for i in range(4)),
        ]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()

        assert errors == []
        assert len(repo.all_words()) == 10
        assert len(repo.all_collections()) == 5
