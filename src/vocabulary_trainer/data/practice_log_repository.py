"""A dedicated, human-readable Excel log of practice sessions (user-requested).

Deliberately separate from two things that already exist and are not this:

* ``master.xlsx`` -- the vocabulary itself. Practice results have never been written
  there (FR-6.16), and this file does not change that.
* ``practice-history.json`` -- the capped, JSON-only figures the in-app summary
  screen persists (score, words practiced, penalties, elapsed time). That store
  answers "how did recent sessions go"; this one answers "when exactly did I
  practice, on what, and which specific words gave me trouble" -- a question best
  answered by something the user can open directly in Excel.

One row per session, appended when a session ends (naturally, by the End session
button, or by closing the window) -- the same three routes that already produce a
:class:`~vocabulary_trainer.domain.models.SessionSummary` (FR-6.12, FR-6.13).
"""

from __future__ import annotations

import os
import threading
import uuid
from datetime import time
from pathlib import Path
from typing import Final

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from vocabulary_trainer.domain.models import PracticeLogEntry

__all__ = ["PracticeLogRepository"]

SHEET_NAME: Final = "Practice Log"
HEADERS: Final = (
    "Date",
    "Start Time",
    "End Time",
    "Collection",
    "Penalty",
    "Finish",
    "Word with penalty",
)
_TIME_NUMBER_FORMAT: Final = "hh:mm:ss"


class PracticeLogRepository:
    """Appends one row per finished-or-ended practice session."""

    def __init__(self, path: Path) -> None:
        self._path = Path(path)
        self._lock = threading.RLock()

    @property
    def path(self) -> Path:
        return self._path

    def append_entry(self, entry: PracticeLogEntry) -> bool:
        """Record one session, reporting success rather than raising.

        Mirrors ``HistoryStore.append``'s posture: the user already finished the
        session and is looking at (or about to see) its summary, so a failure to
        log it here must never surface as an error or block anything. A locked
        file (open in Excel) or any other I/O failure simply means this one row is
        lost, same as a JSON history write failing.
        """
        try:
            with self._lock:
                workbook = self._open_or_create()
                try:
                    self._append_row(workbook[SHEET_NAME], entry)
                    self._save_atomically(workbook)
                finally:
                    workbook.close()
            return True
        except OSError:
            return False

    # ------------------------------------------------------------------

    @staticmethod
    def _append_row(sheet: Worksheet, entry: PracticeLogEntry) -> None:
        row = sheet.max_row + 1
        sheet.cell(row=row, column=1, value=entry.session_date)
        PracticeLogRepository._write_time(sheet, row, 2, entry.start_time)
        PracticeLogRepository._write_time(sheet, row, 3, entry.end_time)
        sheet.cell(row=row, column=4, value=entry.collection_display)
        sheet.cell(row=row, column=5, value=entry.penalty)
        sheet.cell(row=row, column=6, value=entry.finished)
        sheet.cell(row=row, column=7, value=entry.word_with_penalty)

    @staticmethod
    def _write_time(sheet: Worksheet, row: int, column: int, value: time) -> None:
        """Write a time value formatted as ``hh:mm:ss``, not a plain string.

        An actual Excel time cell keeps the column sortable and lets the user do
        duration arithmetic on it directly, which a text column would not.
        """
        cell = sheet.cell(row=row, column=column, value=value)
        cell.number_format = _TIME_NUMBER_FORMAT

    def _open_or_create(self) -> Workbook:
        if not self._path.exists():
            self._path.parent.mkdir(parents=True, exist_ok=True)
            return self._fresh_workbook()

        try:
            workbook = load_workbook(self._path)
        except Exception:
            # A corrupt log file must not block logging a new session -- same
            # "informational, not a system of record" posture as HistoryStore.
            # Starting over loses old rows rather than the new one going unlogged.
            return self._fresh_workbook()

        if SHEET_NAME not in workbook.sheetnames:
            sheet = workbook.create_sheet(SHEET_NAME)
            sheet.append(list(HEADERS))
        return workbook

    @staticmethod
    def _fresh_workbook() -> Workbook:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = SHEET_NAME
        sheet.append(list(HEADERS))
        return workbook

    def _save_atomically(self, workbook: Workbook) -> None:
        """Temp-file-then-swap, same discipline as the master workbook (FR-8.5),
        scaled to this file's lower stakes: a failure here degrades to ``False``
        rather than a typed lock error, because losing one log row is not worth
        the retry machinery a lost vocabulary word would justify.
        """
        temp_path = self._path.parent / f".{self._path.name}.tmp-{uuid.uuid4().hex}"
        try:
            workbook.save(temp_path)
            os.replace(temp_path, self._path)
        except Exception:
            temp_path.unlink(missing_ok=True)
            raise
